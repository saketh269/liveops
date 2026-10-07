"""File connectors, poll mode.

``csv_file``: CSV and Excel (.xlsx) files kept on the Live Ops server under
``LIVEOPS_DATA_DIR/<folder>``. Files arrive through
``POST /api/sources/{id}/upload`` (``app/api/uploads.py``), which sets the
folder to the source id on first upload; an admin can also point ``folder`` at
an existing sub-folder of the data directory. One dataset per file. A file is
re-read only when its size or modification time changes.

``s3_files``: objects in an S3-compatible bucket (AWS S3, MinIO, Ceph, ...)
under a prefix. CSV, XLSX and JSON-lines objects; one dataset per object.
An object is downloaded again only when its ETag changes. boto3 is blocking,
so every call runs in a worker thread, off the event loop.

Safety:
- Folder names and dataset names are checked against what ``discover()``
  lists; nothing outside the data directory can be read (no ``..``, no
  absolute paths, symlinks are skipped).
- Byte caps on every file/object and row caps on every snapshot.
- S3 is read-only (List/Head/Get only); HTTPS is required unless Encryption
  is Off (local testing only).
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.config import get_settings
from app.connectors import netguard
from app.connectors.base import (
    MAX_SNAPSHOT_ROWS,
    Category,
    ConnectorError,
    ConnectorSpec,
    Dataset,
    Mode,
    PollingConnector,
    Record,
    TestReport,
    TestStep,
    check_row_cap,
    normalize_record,
)
from app.connectors.registry import register
from app.connectors.rest import infer_columns, primary_key_guess

MAX_ROWS = MAX_SNAPSHOT_ROWS
XLSX_MAX_UNCOMPRESSED = 200 * 1024 * 1024  # all members together, as declared in the zip
XLSX_MAX_MEMBER = 100 * 1024 * 1024
XLSX_MAX_RATIO = 100  # uncompressed : compressed, for members over XLSX_RATIO_MIN_SIZE
XLSX_RATIO_MIN_SIZE = 1024 * 1024
XLSX_MAX_MEMBERS = 1_000
MAX_COLUMNS = 500  # CSV and XLSX
MAX_CELLS = 2_000_000  # rows x columns, CSV and XLSX (bounds memory, LIVEOPS-71)
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_FILES = 500
DEFAULT_S3_TIMEOUT_S = 30.0
MAX_S3_TIMEOUT_S = 120.0
LOCAL_EXTENSIONS = (".csv", ".xlsx")
S3_EXTENSIONS = (".csv", ".xlsx", ".jsonl", ".ndjson")
SAFE_PART = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")


# --------------------------------------------------------------------------
# Parsing (pure, blocking: call from a worker thread)
# --------------------------------------------------------------------------


def _header(names: list[Any], what: str, *, trim_trailing: bool = False) -> list[str]:
    """Column names from a header row. Linear time (set lookups) and capped (LIVEOPS-71).

    ``trim_trailing`` drops empty cells at the end (xlsx pads rows to the sheet width)."""
    names = list(names)
    while trim_trailing and names and (names[-1] is None or str(names[-1]).strip() == ""):
        names.pop()
    if len(names) > MAX_COLUMNS:
        raise ConnectorError(
            f"{what} has more than {MAX_COLUMNS} columns",
            hint="Keep only the columns you need on the map (ID, status, zone, ...) and upload again.",
        )
    out: list[str] = []
    seen: set[str] = set()
    for i, n in enumerate(names):
        name = str(n).strip() if n is not None else ""
        if not name or name in seen:
            name = f"column_{i + 1}"
            while name in seen:
                name += "_"
        seen.add(name)
        out.append(name)
    return out


def _too_many_cells(what: str) -> ConnectorError:
    return ConnectorError(
        f"{what} has more than {MAX_CELLS:,} cells (rows x columns)",
        hint="Split it into smaller files, or remove columns you don't need.",
    )


def _cap(rows: list[Record], what: str) -> list[Record]:
    """Complete or raise (ADR 0004): parsers stop at MAX_ROWS + 1, then this raises."""
    try:
        check_row_cap(len(rows), what, MAX_ROWS)
    except ConnectorError as e:
        raise ConnectorError(str(e), hint="Split it into smaller files, one per area or site.") from None
    return rows


def check_xlsx_zip(data: bytes, what: str = "The workbook") -> None:
    """Refuse decompression bombs before openpyxl inflates anything (LIVEOPS-19).

    Uses the sizes in the zip directory. Those can't be used to sneak more data
    through: ``zipfile`` stops each member at its declared size and fails the
    CRC check if there is more.
    """
    import zipfile

    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise ConnectorError(
            f"{what} isn't a readable .xlsx file", hint="Save it from Excel as 'Excel Workbook (.xlsx)'."
        ) from None
    too_big = ConnectorError(
        f"{what} expands to too much data to read safely",
        hint="Save a smaller workbook (fewer rows, one sheet), or export it as CSV.",
    )
    if len(infos) > XLSX_MAX_MEMBERS:
        raise too_big
    if sum(i.file_size for i in infos) > XLSX_MAX_UNCOMPRESSED:
        raise too_big
    for i in infos:
        if i.file_size > XLSX_MAX_MEMBER:
            raise too_big
        if i.file_size > XLSX_RATIO_MIN_SIZE and i.file_size > XLSX_MAX_RATIO * max(i.compress_size, 1):
            raise too_big


def parse_csv(data: bytes, delimiter: str = ",", encoding: str = "utf-8-sig", what: str = "The file") -> list[Record]:
    try:
        text = data.decode(encoding)
    except (UnicodeDecodeError, LookupError):
        raise ConnectorError(
            f"{what} isn't valid {encoding} text", hint="Save it as CSV UTF-8, or set the right encoding."
        ) from None
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=(delimiter or ",")[0])
    try:
        head = next(reader, None)
        if head is None:
            return []
        cols = _header(head, what)
        rows: list[Record] = []
        cells = 0
        for line in reader:
            if not any(cell.strip() for cell in line):
                continue
            rec: Record = {c: (line[i] if i < len(line) and line[i] != "" else None) for i, c in enumerate(cols)}
            rows.append(rec)
            cells += len(cols)
            if cells > MAX_CELLS:
                raise _too_many_cells(what)
            if len(rows) > MAX_ROWS:
                break
    except csv.Error as e:
        raise ConnectorError(f"{what} isn't a readable CSV: {e}", hint="Check the delimiter and quoting.") from None
    return _cap(rows, what)


def parse_xlsx(data: bytes, sheet: str | None = None, what: str = "The workbook") -> list[Record]:
    import zipfile

    from openpyxl import load_workbook

    check_xlsx_zip(data, what)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise ConnectorError(
            f"{what} isn't a readable .xlsx file", hint="Save it from Excel as 'Excel Workbook (.xlsx)'."
        ) from None
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                raise ConnectorError(
                    f"{what} has no sheet named {sheet!r}", hint=f"Sheets in it: {', '.join(wb.sheetnames)[:200]}"
                )
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        # max_col: a sheet can claim 16,384 columns; never build rows that wide.
        it = ws.iter_rows(values_only=True, max_col=MAX_COLUMNS + 1)
        head = next(it, None)
        if head is None:
            return []
        cols = _header(list(head), what, trim_trailing=True)
        rows: list[Record] = []
        cells = 0
        for line in it:
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in line):
                continue
            rows.append(normalize_record({c: (line[i] if i < len(line) else None) for i, c in enumerate(cols)}))
            cells += len(cols)
            if cells > MAX_CELLS:
                raise _too_many_cells(what)
            if len(rows) > MAX_ROWS:
                break
        return _cap(rows, what)
    finally:
        wb.close()


def parse_jsonl(data: bytes, what: str = "The file") -> list[Record]:
    rows: list[Record] = []
    for n, line in enumerate(data.decode("utf-8-sig", errors="replace").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            obj = json.loads(line)
        except ValueError:
            raise ConnectorError(
                f"{what} line {n} isn't valid JSON", hint="JSON-lines files need one JSON object per line."
            ) from None
        if isinstance(obj, dict):
            rows.append(normalize_record(obj))
        if len(rows) > MAX_ROWS:
            break
    return _cap(rows, what)


def parse_bytes(name: str, data: bytes, settings: dict[str, Any]) -> list[Record]:
    ext = os.path.splitext(name)[1].lower()
    what = f"File {name!r}"
    if ext == ".csv":
        return parse_csv(data, settings.get("delimiter") or ",", settings.get("encoding") or "utf-8-sig", what)
    if ext == ".xlsx":
        return parse_xlsx(data, settings.get("sheet") or None, what)
    if ext in (".jsonl", ".ndjson"):
        return parse_jsonl(data, what)
    raise ConnectorError(f"{what} has an unsupported type", hint="Use .csv, .xlsx or .jsonl.")


def _dataset(name: str, rows: list[Record]) -> Dataset:
    cols = infer_columns(rows)
    return Dataset(name=name, columns=cols, primary_key=primary_key_guess(cols))


def _too_big(name: str, limit: int = MAX_FILE_BYTES) -> ConnectorError:
    return ConnectorError(
        f"File {name!r} is larger than {limit // (1024 * 1024)} MB", hint="Split it into smaller files."
    )


# --------------------------------------------------------------------------
# Local files (uploaded)
# --------------------------------------------------------------------------


def data_dir() -> Path:
    return Path(get_settings().data_dir).resolve()


def safe_folder(folder: str) -> Path:
    """Resolve a folder setting to a directory inside the data dir, or refuse."""
    parts = [p for p in str(folder).replace("\\", "/").split("/") if p]
    if not parts or str(folder).startswith(("/", "\\")) or not all(SAFE_PART.match(p) for p in parts):
        raise ConnectorError(
            f"Folder {folder!r} isn't allowed",
            hint="Use a simple folder name inside the Live Ops data directory (letters, digits, . _ -), "
            "with no '..' or leading '/'.",
        )
    root = data_dir()
    path = root.joinpath(*parts).resolve()
    if not path.is_relative_to(root):
        raise ConnectorError(f"Folder {folder!r} is outside the data directory", hint="Pick a folder inside it.")
    return path


@dataclass
class _Cached:
    sig: tuple[Any, ...]
    rows: list[Record]


@register
class CsvFileConnector(PollingConnector):
    spec = ConnectorSpec(
        type="csv_file",
        display_name="CSV / Excel file",
        category=Category.FILE,
        modes=[Mode.POLL],
        description="Upload CSV or Excel (.xlsx) files. Upload a new version to update the map.",
        maturity="beta",
        settings_schema={
            "type": "object",
            "properties": {
                "folder": {
                    "type": "string",
                    "title": "Folder on the server",
                    "description": "Leave empty: uploads go to this source's own folder. Admins can point it "
                    "at another folder inside LIVEOPS_DATA_DIR.",
                },
                "delimiter": {"type": "string", "title": "CSV delimiter", "default": ",", "maxLength": 1},
                "encoding": {"type": "string", "title": "CSV encoding", "default": "utf-8-sig"},
                "sheet": {"type": "string", "title": "Excel sheet (default: first sheet)"},
            },
        },
    )

    def __init__(self, settings: dict[str, Any], secrets: dict[str, Any], *, source_id: str | None = None) -> None:
        super().__init__(settings, secrets, source_id=source_id)
        self._cache: dict[str, _Cached] = {}

    def folder_name(self) -> str:
        """``folder`` setting, else the source's own folder named by its id (ADR 0004)."""
        f = self.settings.get("folder") or self.source_id
        if not f:
            raise ConnectorError("This source has no folder yet", hint="Save the source, then upload a file to it.")
        return str(f)

    def folder(self) -> Path:
        return safe_folder(self.folder_name())

    def _list(self) -> list[Path]:
        root = self.folder()
        if not root.is_dir():
            raise ConnectorError(
                "No file has been uploaded to this source yet",
                hint="Upload a .csv or .xlsx file to this source (or check the folder setting).",
            )
        out = []
        for p in sorted(root.iterdir()):
            if p.name.startswith(".") or p.is_symlink() or not p.is_file():
                continue
            if p.suffix.lower() in LOCAL_EXTENSIONS:
                out.append(p)
            if len(out) >= MAX_FILES:
                break
        return out

    def _path(self, dataset: str) -> Path:
        for p in self._list():
            if p.name == dataset:
                return p
        raise ConnectorError(
            f"File {dataset!r} isn't in this source", hint="Pick one of the files listed, or upload it first."
        )

    def _read(self, path: Path) -> list[Record]:
        st = path.stat()
        sig = (st.st_ino, st.st_mtime_ns, st.st_size)
        cached = self._cache.get(path.name)
        if cached is not None and cached.sig == sig:
            return cached.rows
        if st.st_size > MAX_FILE_BYTES:
            raise _too_big(path.name)
        with path.open("rb") as fh:
            data = fh.read(MAX_FILE_BYTES + 1)
        rows = parse_bytes(path.name, data, self.settings)
        self._cache[path.name] = _Cached(sig, rows)
        return rows

    async def test(self) -> TestReport:
        started = time.monotonic()
        steps: list[TestStep] = []
        try:
            files = await asyncio.to_thread(self._list)
            steps.append(TestStep(name="Data folder", ok=True, detail=self.folder_name()))
        except ConnectorError as e:
            steps.append(TestStep(name="Data folder", ok=False, detail=str(e), hint=e.hint))
            return TestReport.from_steps(steps, started)
        steps.append(
            TestStep(
                name="Files",
                ok=bool(files),
                detail=f"{len(files)} .csv/.xlsx files",
                hint="" if files else "Upload a .csv or .xlsx file to this source.",
            )
        )
        for p in files[:5]:
            try:
                rows = await asyncio.to_thread(self._read, p)
                steps.append(TestStep(name=f"Read {p.name}", ok=True, detail=f"{len(rows)} rows"))
            except ConnectorError as e:
                steps.append(TestStep(name=f"Read {p.name}", ok=False, detail=str(e), hint=e.hint))
        return TestReport.from_steps(steps, started)

    async def discover(self) -> list[Dataset]:
        def run() -> list[Dataset]:
            return [_dataset(p.name, self._read(p)) for p in self._list()]

        return await asyncio.to_thread(run)

    async def snapshot(self, dataset: str) -> list[Record]:
        return await asyncio.to_thread(lambda: self._read(self._path(dataset)))


# --------------------------------------------------------------------------
# S3-compatible object storage
# --------------------------------------------------------------------------


@register
class S3FilesConnector(PollingConnector):
    spec = ConnectorSpec(
        type="s3_files",
        display_name="S3 / object storage files",
        category=Category.FILE,
        modes=[Mode.POLL],
        description="Reads CSV, Excel and JSON-lines files from an S3-compatible bucket (AWS S3, MinIO, ...).",
        maturity="needs_real_test",
        settings_schema={
            "type": "object",
            "required": ["bucket"],
            "properties": {
                "bucket": {"type": "string", "title": "Bucket"},
                "prefix": {"type": "string", "title": "Folder / prefix", "default": "", "examples": ["exports/"]},
                "region": {"type": "string", "title": "Region", "default": "us-east-1"},
                "endpoint_url": {
                    "type": "string",
                    "title": "Endpoint URL (not AWS)",
                    "description": "For MinIO and other S3-compatible stores, e.g. https://minio.example.com:9000. "
                    "Leave empty for AWS.",
                },
                "encryption": {
                    "type": "string",
                    "title": "Encryption",
                    "enum": ["required", "verify", "off"],
                    "default": "required",
                    "description": "HTTPS with certificate checks. Use 'off' (plain HTTP) only for local testing.",
                },
                "allow_private_network": {
                    "type": "boolean",
                    "title": "Allow private network addresses",
                    "default": False,
                    "description": "Needed for an endpoint inside your company network (10.x, 192.168.x, "
                    "localhost). Cloud metadata and link-local addresses are always blocked.",
                },
                "timeout_s": {
                    "type": "number",
                    "title": "Timeout (seconds)",
                    "default": 30,
                    "minimum": 1,
                    "maximum": 120,
                },
                "use_instance_role": {
                    "type": "boolean",
                    "title": "Use the Live Ops server's own AWS role (admin only)",
                    "default": False,
                    "description": "Uses the server's IAM role instead of an access key. Works only if the "
                    "administrator set LIVEOPS_S3_ALLOW_INSTANCE_ROLE=true. Leave off.",
                },
                "delimiter": {"type": "string", "title": "CSV delimiter", "default": ",", "maxLength": 1},
                "encoding": {"type": "string", "title": "CSV encoding", "default": "utf-8-sig"},
                "sheet": {"type": "string", "title": "Excel sheet (default: first sheet)"},
            },
        },
        secrets_schema={
            "type": "object",
            "properties": {
                "access_key_id": {"type": "string", "title": "Access key ID"},
                "secret_access_key": {"type": "string", "title": "Secret access key", "format": "password"},
                "session_token": {"type": "string", "title": "Session token (optional)", "format": "password"},
            },
        },
    )

    def __init__(self, settings: dict[str, Any], secrets: dict[str, Any], *, source_id: str | None = None) -> None:
        super().__init__(settings, secrets, source_id=source_id)
        self._client: Any = None
        self._cache: dict[str, _Cached] = {}
        self._keys: set[str] | None = None
        self._pin = netguard.PinnedAddress()

    @property
    def bucket(self) -> str:
        return str(self.settings.get("bucket") or "")

    @property
    def prefix(self) -> str:
        return str(self.settings.get("prefix") or "")

    def _endpoint(self) -> tuple[str, str, int]:
        """(endpoint URL, host, port). Always explicit, so nothing from the server's
        AWS config or env (AWS_ENDPOINT_URL*, ~/.aws/config) can redirect it (LIVEOPS-66)."""
        from urllib.parse import urlsplit

        region = str(self.settings.get("region") or "us-east-1")
        if not re.fullmatch(r"[a-z0-9-]{1,32}", region):
            raise ConnectorError("Region must look like us-east-1", hint="Check the region setting.")
        endpoint = str(self.settings.get("endpoint_url") or "").strip() or f"https://s3.{region}.amazonaws.com"
        parts = urlsplit(endpoint)
        scheme = parts.scheme.lower()
        if scheme not in ("http", "https"):
            raise ConnectorError(
                "Endpoint URL must start with https://",
                hint="Use the full address, e.g. https://minio.example.com:9000.",
            )
        if scheme == "http" and (self.settings.get("encryption") or "required") != "off":
            raise ConnectorError(
                "The endpoint uses plain http://, which isn't encrypted",
                hint="Use https://, or set Encryption to Off for local testing only.",
            )
        if not parts.hostname or parts.path not in ("", "/") or parts.query:
            raise ConnectorError(
                "The endpoint URL must be just a scheme, host and port", hint="Use e.g. https://minio.example.com:9000."
            )
        return endpoint, parts.hostname, parts.port or (443 if scheme == "https" else 80)

    def _timeout(self) -> float:
        try:
            t = float(self.settings.get("timeout_s") or DEFAULT_S3_TIMEOUT_S)
        except (TypeError, ValueError):
            t = DEFAULT_S3_TIMEOUT_S
        return max(1.0, min(t, MAX_S3_TIMEOUT_S))

    def _s3(self) -> Any:
        if self._client is not None:
            return self._client
        import botocore.session
        from botocore.config import Config

        endpoint, _, _ = self._endpoint()
        creds = self._credentials()
        # An isolated botocore session: no AWS_* env vars, no ~/.aws files, no
        # configured endpoints, no env proxies (LIVEOPS-21/66).
        session_vars = {k: (v[0], None, v[2], v[3]) for k, v in botocore.session.Session.SESSION_VARIABLES.items()}
        session = botocore.session.Session(session_vars=session_vars)
        session.set_config_variable("config_file", os.devnull)
        session.set_config_variable("credentials_file", os.devnull)
        session.set_config_variable("ignore_configured_endpoint_urls", True)
        if not creds:  # admin-approved instance role: borrow the default credential chain only
            session._credentials = botocore.session.Session().get_credentials()
        timeout = self._timeout()
        cfg = Config(
            connect_timeout=min(timeout, 10.0),
            read_timeout=timeout,
            retries={"max_attempts": 2, "mode": "standard"},
            s3={"addressing_style": "path"},  # the host we checked is the host we call
            proxies={},
            user_agent_extra="liveops",
        )
        client = session.create_client(
            "s3",
            endpoint_url=endpoint,
            region_name=str(self.settings.get("region") or "us-east-1"),
            **creds,
            config=cfg,
            verify=True,  # https always verifies; 'off' only permits plain http:// (LIVEOPS-66)
        )
        netguard.pin_botocore_client(client, self._pin)
        self._client = client
        return client

    def _credentials(self) -> dict[str, str]:
        """Only the keys entered on this source. Never boto3's default chain
        (env vars, ~/.aws, instance role) unless the admin allows it and the
        source asks for it (LIVEOPS-16)."""
        key_id = str(self.secrets.get("access_key_id") or "").strip()
        secret = str(self.secrets.get("secret_access_key") or "")
        if key_id and secret:
            out = {"aws_access_key_id": key_id, "aws_secret_access_key": secret}
            if self.secrets.get("session_token"):
                out["aws_session_token"] = str(self.secrets["session_token"])
            return out
        if self.settings.get("use_instance_role") is True:
            if get_settings().s3_allow_instance_role:
                return {}  # explicit, admin-approved opt-in: boto3's default chain
            raise ConnectorError(
                "Using the server's own AWS role is turned off on this server",
                hint="Enter an access key ID and secret access key, or ask the administrator to set "
                "LIVEOPS_S3_ALLOW_INSTANCE_ROLE=true.",
            )
        raise ConnectorError(
            "Enter an access key ID and secret access key",
            hint="Create a read-only key (s3:ListBucket and s3:GetObject) for this bucket and enter both values.",
        )

    def _check_endpoint(self) -> None:
        """Apply the outbound network policy to the endpoint (AWS too), on every call,
        and pin the client's new connections to the IP that passed (LIVEOPS-21)."""
        _, host, port = self._endpoint()
        ip = netguard.check_host_sync(host, port, self.settings.get("allow_private_network") is True)
        self._pin.ip = str(ip)

    async def close(self) -> None:
        client, self._client = self._client, None
        if client is not None:
            await asyncio.to_thread(client.close)

    async def _run(self, fn: Any, *args: Any) -> Any:
        deadline = 3 * self._timeout() + 10  # overall per call: retries can't stretch a poll forever
        try:
            async with asyncio.timeout(deadline):
                return await asyncio.to_thread(fn, *args)
        except TimeoutError:
            raise ConnectorError(
                f"The storage service took longer than {deadline:g} s",
                hint="Check the endpoint and network, or use smaller files.",
            ) from None
        except ConnectorError:
            raise
        except Exception as e:  # noqa: BLE001 - translated into a plain-English error
            raise _s3_error(e, self.bucket) from None

    def _list_keys(self) -> list[tuple[str, int]]:
        self._check_endpoint()
        out: list[tuple[str, int]] = []
        paginator = self._s3().get_paginator("list_objects_v2")
        for page in paginator.paginate(Bucket=self.bucket, Prefix=self.prefix, PaginationConfig={"PageSize": 1000}):
            for obj in page.get("Contents", []):
                key = obj["Key"]
                if key.lower().endswith(S3_EXTENSIONS):
                    out.append((key, int(obj.get("Size", 0))))
                if len(out) >= MAX_FILES:
                    return out
        return out

    def _read(self, key: str) -> list[Record]:
        self._check_endpoint()
        s3 = self._s3()
        head = s3.head_object(Bucket=self.bucket, Key=key)
        sig = (head.get("ETag"), head.get("ContentLength"))
        cached = self._cache.get(key)
        if cached is not None and cached.sig == sig:
            return cached.rows
        if int(head.get("ContentLength") or 0) > MAX_FILE_BYTES:
            raise _too_big(key)
        # IfMatch: fail instead of mixing two versions if the object changes between head and get.
        extra = {"IfMatch": head["ETag"]} if head.get("ETag") else {}
        obj = s3.get_object(Bucket=self.bucket, Key=key, **extra)
        body = obj["Body"]
        try:
            data = body.read(MAX_FILE_BYTES + 1)
        finally:
            body.close()
        if len(data) > MAX_FILE_BYTES:
            raise _too_big(key)
        rows = parse_bytes(key, data, self.settings)
        self._cache[key] = _Cached((obj.get("ETag") or head.get("ETag"), len(data)), rows)
        return rows

    async def _resolve(self, dataset: str) -> str:
        if self._keys is None or dataset not in self._keys:
            self._keys = {k for k, _ in await self._run(self._list_keys)}
        if dataset not in self._keys:
            raise ConnectorError(
                f"Object {dataset!r} isn't a readable file under this prefix",
                hint="Pick one of the files listed. Supported: .csv, .xlsx, .jsonl.",
            )
        return dataset

    async def test(self) -> TestReport:
        started = time.monotonic()
        steps: list[TestStep] = []
        try:
            await self._run(self._s3)
            endpoint, _, _ = self._endpoint()
            https = endpoint.lower().startswith("https://")
            steps.append(
                TestStep(
                    name="Settings, encryption and access key",
                    ok=True,
                    detail=f"{endpoint}: "
                    + ("HTTPS, certificate checked" if https else "plain HTTP (local testing only)"),
                )
            )
        except ConnectorError as e:
            steps.append(TestStep(name="Settings, encryption and access key", ok=False, detail=str(e), hint=e.hint))
            return TestReport.from_steps(steps, started)
        try:
            keys = await self._run(self._list_keys)
            steps.append(TestStep(name="Sign in and list the bucket", ok=True, detail=f"bucket {self.bucket!r}"))
        except ConnectorError as e:
            steps.append(TestStep(name="Sign in and list the bucket", ok=False, detail=str(e), hint=e.hint))
            return TestReport.from_steps(steps, started)
        steps.append(
            TestStep(
                name="Files",
                ok=bool(keys),
                detail=f"{len(keys)} .csv/.xlsx/.jsonl files under {self.prefix or 'the bucket root'!r}",
                hint="" if keys else "Check the prefix, and that files end in .csv, .xlsx or .jsonl.",
            )
        )
        if keys:
            try:
                rows = await self._run(self._read, keys[0][0])
                steps.append(TestStep(name="Read a file", ok=True, detail=f"{keys[0][0]}: {len(rows)} rows"))
            except ConnectorError as e:
                steps.append(TestStep(name="Read a file", ok=False, detail=str(e), hint=e.hint))
        return TestReport.from_steps(steps, started)

    async def discover(self) -> list[Dataset]:
        keys = await self._run(self._list_keys)
        self._keys = {k for k, _ in keys}
        out = []
        for k, _ in keys:
            out.append(_dataset(k, await self._run(self._read, k)))
        return out

    async def snapshot(self, dataset: str) -> list[Record]:
        key = await self._resolve(dataset)
        return await self._run(self._read, key)


def _s3_error(e: Exception, bucket: str) -> ConnectorError:
    from botocore.exceptions import ClientError, EndpointConnectionError, NoCredentialsError

    if isinstance(e, ClientError):
        code = str(e.response.get("Error", {}).get("Code", ""))
        if code in ("InvalidAccessKeyId", "SignatureDoesNotMatch", "InvalidToken", "ExpiredToken", "403"):
            return ConnectorError(
                f"The storage service refused the access key ({code})",
                hint="Check the access key ID and secret access key.",
            )
        if code in ("AccessDenied", "AllAccessDisabled"):
            return ConnectorError(
                f"The access key may not read bucket {bucket!r} (AccessDenied)",
                hint="Give the key s3:ListBucket on the bucket and s3:GetObject on the prefix.",
            )
        if code in ("NoSuchBucket", "404", "NotFound"):
            return ConnectorError(f"Bucket or file not found in {bucket!r} ({code})", hint="Check the bucket name.")
        if code == "PreconditionFailed":
            return ConnectorError("The file changed while it was being read", hint="It will be read again next poll.")
        return ConnectorError(f"The storage service answered {code or 'an error'}", hint="Check bucket and prefix.")
    if isinstance(e, NoCredentialsError):
        return ConnectorError("No access key is set", hint="Enter the access key ID and secret access key.")
    if isinstance(e, EndpointConnectionError):
        return ConnectorError(
            "Couldn't reach the storage endpoint",
            hint="Check the endpoint URL and region, and that a firewall allows the connection.",
        )
    return ConnectorError(f"Storage error: {type(e).__name__}", hint="Check the endpoint URL and network.")
